#!/usr/bin/env python3
"""Generate the synthetic dataset for the SMBC Data & AI Workshop.

The data describes a made-up SMBC APAC corporate lending book: counterparty
groups and their subsidiaries, credit facilities, month-end balances, FX rates,
and the risk team's spreadsheet of group limits. Every name and number is
invented.

Two things are planted on purpose and recorded in facilitator/answer_key.json:
  * data-quality problems that Lab 1 has to fix, and
  * findings that Lab 2, the Genie Code segment and the quiz depend on.

Only the Python standard library is required (openpyxl is used for the .xlsx
when installed), so the same script runs on a laptop or in a Databricks notebook:

    python generate_data.py --out ./output
    python generate_data.py --out /Volumes/smbc_workshop/raw

Files land in <out>/landing/... (for participants) and <out>/facilitator/...
(answer key and a CSV copy of the limits; keep these away from participants).
"""

import argparse
import csv
import datetime as dt
import json
import math
import os
import random
import shutil
import tempfile
import zipfile
from collections import defaultdict
from xml.sax.saxutils import escape

SEED = 20261022

# ---------------------------------------------------------------------------
# Calendar
# ---------------------------------------------------------------------------


def month_ends(first_year, first_month, count):
    """Return `count` consecutive month-end dates, starting with the given month."""
    dates, year, month = [], first_year, first_month
    for _ in range(count):
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
        dates.append(dt.date(year, month, 1) - dt.timedelta(days=1))
    return dates


# April 2025 (start of FY2025) to September 2026 (end of FY2026 H1).
MONTHS = month_ends(2025, 4, 18)
APR26 = MONTHS.index(dt.date(2026, 4, 30))
JUN26 = MONTHS.index(dt.date(2026, 6, 30))
SEP26 = MONTHS.index(dt.date(2026, 9, 30))


def add_months(d, months):
    years, month0 = divmod(d.month - 1 + months, 12)
    year, month = d.year + years, month0 + 1
    next_first = dt.date(year + 1, 1, 1) if month == 12 else dt.date(year, month + 1, 1)
    return dt.date(year, month, min(d.day, (next_first - dt.timedelta(days=1)).day))


def months_between(a, b):
    return (b.year - a.year) * 12 + (b.month - a.month)


def last_business_day(d):
    while d.weekday() >= 5:
        d -= dt.timedelta(days=1)
    return d


def fmt_date(d, source_system):
    # The trade-finance system exports DD/MM/YYYY; the core loan system exports ISO dates.
    return d.strftime("%d/%m/%Y") if source_system == "TRADE_FIN" else d.isoformat()


# ---------------------------------------------------------------------------
# Reference data
# ---------------------------------------------------------------------------

CURRENCIES = {
    # code: (USD per unit at end-April 2025, monthly volatility)
    "USD": (1.0, 0.0),
    "JPY": (0.00670, 0.015),
    "SGD": (0.7600, 0.006),
    "HKD": (0.12880, 0.0005),
    "CNY": (0.13800, 0.004),
    "THB": (0.02950, 0.010),
    "VND": (0.0000392, 0.004),
    "IDR": (0.0000605, 0.008),
    "INR": (0.01170, 0.006),
    "AUD": (0.6400, 0.012),
    "KRW": (0.000720, 0.010),
    "MYR": (0.2250, 0.007),
    "PHP": (0.01760, 0.006),
    "TWD": (0.03100, 0.006),
}

# Planted: these currencies strengthen against USD between end-April and end-September
# 2026, lifting the USD value of Orchid Group's local-currency borrowing. The yen
# weakens a little over the same period.
PLANTED_FX_MOVE = {"THB": 0.10, "IDR": 0.09, "VND": 0.05, "INR": 0.04, "SGD": 0.05, "JPY": -0.03}

COUNTRIES = {
    # code: (name, currency, home booking branch, legal suffix, local script)
    "SG": ("Singapore", "SGD", "SIN", "Pte. Ltd.", None),
    "HK": ("Hong Kong", "HKD", "HKG", "Limited", "zht"),
    "TH": ("Thailand", "THB", "BKK", "Co., Ltd.", "th"),
    "VN": ("Vietnam", "VND", "SGN", "JSC", "vi"),
    "ID": ("Indonesia", "IDR", "JKT", "", None),
    "IN": ("India", "INR", "BOM", "Private Limited", None),
    "AU": ("Australia", "AUD", "SYD", "Pty Ltd", None),
    "JP": ("Japan", "JPY", "SIN", "K.K.", "ja"),
    "KR": ("South Korea", "KRW", "SEL", "Co., Ltd.", "ko"),
    "CN": ("China", "CNY", "SHA", "Co., Ltd.", "zh"),
    "MY": ("Malaysia", "MYR", "KUL", "Sdn. Bhd.", None),
    "PH": ("Philippines", "PHP", "MNL", "Inc.", None),
    "TW": ("Taiwan", "TWD", "TPE", "Co., Ltd.", "zht"),
}
COUNTRY_WEIGHTS = {"SG": 16, "HK": 12, "TH": 9, "VN": 7, "ID": 9, "IN": 11, "AU": 9,
                   "JP": 5, "KR": 5, "CN": 6, "MY": 5, "PH": 3, "TW": 3}

SECTORS = {
    # sector: (word used in English names, Indonesian word, weight, product mix)
    "Real Estate": ("Land", "Properti", 9, "property"),
    "Shipping & Logistics": ("Logistics", "Logistik", 7, "corporate"),
    "Energy": ("Energy", "Energi", 7, "project"),
    "Utilities": ("Power", "Daya", 5, "project"),
    "Manufacturing": ("Industries", "Industri", 11, "trade"),
    "Technology": ("Technologies", "Teknologi", 7, "corporate"),
    "Telecommunications": ("Telecom", "Telekomunikasi", 4, "corporate"),
    "Consumer Goods": ("Consumer", "Konsumen", 7, "trade"),
    "Healthcare": ("Healthcare", "Kesehatan", 4, "corporate"),
    "Financial Institutions": ("Capital", "Kapital", 6, "corporate"),
    "Infrastructure": ("Infrastructure", "Infrastruktur", 5, "project"),
    "Agriculture": ("Agri", "Agri", 4, "trade"),
    "Metals & Mining": ("Resources", "Tambang", 4, "trade"),
    "Chemicals": ("Chemicals", "Kimia", 5, "trade"),
    "Transportation": ("Transport", "Transportasi", 5, "corporate"),
    "Retail": ("Retail", "Ritel", 5, "trade"),
}

SECTOR_LOCAL = {
    # sector: (ja, zh-Hans, zh-Hant, th, ko, vi)
    "Real Estate": ("不動産", "置业", "置業", "พร็อพเพอร์ตี้", "부동산", "Bất động sản"),
    "Shipping & Logistics": ("ロジスティクス", "物流", "物流", "โลจิสติกส์", "로지스틱스", "Logistics"),
    "Energy": ("エナジー", "能源", "能源", "เอ็นเนอร์ยี", "에너지", "Năng lượng"),
    "Utilities": ("パワー", "电力", "電力", "พาวเวอร์", "파워", "Điện lực"),
    "Manufacturing": ("工業", "工业", "工業", "อินดัสทรีส์", "산업", "Công nghiệp"),
    "Technology": ("テクノロジーズ", "科技", "科技", "เทคโนโลยี", "테크놀로지", "Công nghệ"),
    "Telecommunications": ("テレコム", "电讯", "電訊", "เทเลคอม", "텔레콤", "Viễn thông"),
    "Consumer Goods": ("コンシューマー", "消费品", "消費品", "คอนซูมเมอร์", "컨슈머", "Hàng tiêu dùng"),
    "Healthcare": ("ヘルスケア", "医疗", "醫療", "เฮลท์แคร์", "헬스케어", "Y tế"),
    "Financial Institutions": ("キャピタル", "资本", "資本", "แคปปิตอล", "캐피탈", "Tài chính"),
    "Infrastructure": ("インフラ", "基建", "基建", "อินฟราสตรัคเจอร์", "인프라", "Hạ tầng"),
    "Agriculture": ("アグリ", "农业", "農業", "อะกริ", "애그리", "Nông nghiệp"),
    "Metals & Mining": ("リソーシズ", "矿业", "礦業", "รีซอร์สเซส", "리소시스", "Khoáng sản"),
    "Chemicals": ("ケミカル", "化工", "化工", "เคมิคอล", "케미칼", "Hóa chất"),
    "Transportation": ("トランスポート", "运输", "運輸", "ทรานสปอร์ต", "트랜스포트", "Vận tải"),
    "Retail": ("リテール", "零售", "零售", "รีเทล", "리테일", "Bán lẻ"),
}
SCRIPT_INDEX = {"ja": 0, "zh": 1, "zht": 2, "th": 3, "ko": 4, "vi": 5}

# Name stems that have local-script versions, used for groups based in HK, TH, VN, JP, KR, CN, TW.
LOCAL_STEMS = {
    # English: (ja, zh-Hans, zh-Hant, th, ko, vi)
    "Golden Harbour": ("ゴールデン・ハーバー", "金港", "金港", "โกลเด้น ฮาร์เบอร์", "골든 하버", "Cảng Vàng"),
    "Jade River": ("ジェイド・リバー", "翡翠河", "翡翠河", "เจด ริเวอร์", "제이드 리버", "Sông Ngọc"),
    "Silver Peak": ("シルバー・ピーク", "银峰", "銀峰", "ซิลเวอร์ พีค", "실버 피크", "Đỉnh Bạc"),
    "Red Phoenix": ("レッド・フェニックス", "红凤凰", "紅鳳凰", "เรด ฟีนิกซ์", "레드 피닉스", "Phượng Hoàng Đỏ"),
    "Eastern Star": ("イースタン・スター", "东方之星", "東方之星", "อีสเทิร์น สตาร์", "이스턴 스타", "Sao Phương Đông"),
    "Blue Ocean": ("ブルー・オーシャン", "蓝海", "藍海", "บลู โอเชียน", "블루 오션", "Đại Dương Xanh"),
    "Sunrise": ("サンライズ", "旭日", "旭日", "ซันไรส์", "선라이즈", "Bình Minh"),
    "Bamboo Grove": ("バンブー・グローブ", "竹林", "竹林", "แบมบู โกรฟ", "뱀부 그로브", "Rừng Tre"),
    "Coral Bay": ("コーラル・ベイ", "珊瑚湾", "珊瑚灣", "คอรัล เบย์", "코랄 베이", "Vịnh San Hô"),
    "Twin Dragon": ("ツイン・ドラゴン", "双龙", "雙龍", "ทวิน ดราก้อน", "트윈 드래곤", "Song Long"),
    "White Crane": ("ホワイト・クレイン", "白鹤", "白鶴", "ไวท์ เครน", "화이트 크레인", "Hạc Trắng"),
    "Mountain Spring": ("マウンテン・スプリング", "山泉", "山泉", "เมาน์เทน สปริง", "마운틴 스프링", "Suối Núi"),
    "Lucky Cloud": ("ラッキー・クラウド", "祥云", "祥雲", "ลัคกี้ คลาวด์", "럭키 클라우드", "Mây Lành"),
    "Pine Ridge": ("パイン・リッジ", "松岭", "松嶺", "ไพน์ ริดจ์", "파인 리지", "Đồi Thông"),
    "Morning Tide": ("モーニング・タイド", "晨潮", "晨潮", "มอร์นิ่ง ไทด์", "모닝 타이드", "Triều Sớm"),
    "Harmony": ("ハーモニー", "和谐", "和諧", "ฮาร์โมนี", "하모니", "Hòa Hợp"),
    "Amber Gate": ("アンバー・ゲート", "琥珀门", "琥珀門", "แอมเบอร์ เกต", "앰버 게이트", "Cổng Hổ Phách"),
    "Northern Light": ("ノーザン・ライト", "北极光", "北極光", "นอร์เทิร์น ไลท์", "노던 라이트", "Bắc Quang"),
}
OTHER_STEMS = [
    "Aurora", "Banyan", "Cedar", "Crescent", "Delta Bay", "Emerald", "Falcon", "Horizon", "Ironbark",
    "Kestrel", "Mangrove", "Monsoon", "Nova", "Oakridge", "Pearl Strait", "Quantum", "Redwood",
    "Sapphire", "Southern Cross", "Summit Ridge", "Teakwood", "Tiger Bay", "Unity", "Vista",
    "Willow Creek", "Zenith", "Lion Gate", "Mekong Valley", "Andaman", "Java Sea", "Indus", "Han River",
    "Straits", "Monarch", "Atlas", "Orion", "Pinnacle", "Kirin", "Garuda", "Naga", "Saffron", "Cobalt",
    "Granite", "Silverleaf", "Brightwater", "Stonebridge", "Rivergate", "Highland", "Seabreeze", "Merlion",
]
GROUP_SUFFIXES = ["Holdings", "Group", "Corporation", "Industries", "Capital", "Partners",
                  "Enterprises", "International"]
NAME_QUALIFIERS = ["Trading", "Services", "Ventures", "Solutions", "Asia", "Development", "Pacific", "Global",
                   "Regional", "Investments", "Operations", "Management", "Systems", "Projects", "Network",
                   "Supply", "Distribution", "Engineering", "Assets", "Partners", "Enterprises", "Holdings"]

PRODUCTS = {
    # product: (source system, min tenor months, max tenor months, balance profile, expected utilisation)
    "Term Loan": ("CORE_LOANS", 36, 84, "amortising", 0.70),
    "Revolving Credit Facility": ("CORE_LOANS", 12, 60, "revolving", 0.55),
    "Working Capital Loan": ("CORE_LOANS", 12, 36, "revolving", 0.55),
    "Project Finance Loan": ("CORE_LOANS", 96, 180, "construction", 0.70),
    "Trade Finance - Import LC": ("TRADE_FIN", 12, 24, "trade", 0.60),
    "Trade Finance - Export Loan": ("TRADE_FIN", 12, 24, "trade", 0.60),
}
PRODUCT_MIX = {
    "trade": {"Trade Finance - Import LC": 25, "Trade Finance - Export Loan": 20,
              "Working Capital Loan": 20, "Term Loan": 20, "Revolving Credit Facility": 15},
    "project": {"Project Finance Loan": 40, "Term Loan": 35, "Revolving Credit Facility": 15,
                "Trade Finance - Import LC": 10},
    "property": {"Term Loan": 65, "Revolving Credit Facility": 20, "Working Capital Loan": 15},
    "corporate": {"Term Loan": 35, "Revolving Credit Facility": 35, "Working Capital Loan": 15,
                  "Trade Finance - Import LC": 15},
}
COLLATERAL = {
    "Real Estate": "Real Estate", "Shipping & Logistics": "Vessel", "Energy": "Plant & Equipment",
    "Utilities": "Plant & Equipment", "Infrastructure": "Plant & Equipment",
    "Metals & Mining": "Plant & Equipment", "Financial Institutions": "Cash",
}

# Illustrative only; not real emission factors.
EMISSION_FACTORS = {
    "Real Estate": 95, "Shipping & Logistics": 480, "Energy": 910, "Utilities": 1240,
    "Manufacturing": 360, "Technology": 35, "Telecommunications": 55, "Consumer Goods": 140,
    "Healthcare": 60, "Financial Institutions": 15, "Infrastructure": 260, "Agriculture": 520,
    "Metals & Mining": 870, "Chemicals": 640, "Transportation": 430, "Retail": 75,
}

CREDIT_OFFICERS = ["A. Tan", "R. Kumar", "M. Sato", "J. Lim", "S. Wong", "P. Nguyen", "K. Lee",
                   "D. Santos", "N. Sharma", "H. Wijaya", "L. Chen", "T. Suzuki"]

# ---------------------------------------------------------------------------
# Portfolio shape
# ---------------------------------------------------------------------------

N_GROUPS = 420
N_BIG_TICKET = 14          # single-name groups big enough that Orchid's subsidiaries never top the counterparty list
GROUP_CAP_USD = 1.08e9     # no ordinary group goes above this in any month
DIRTY_UTILISATION_CAP = 0.88  # even with nothing cleaned, no other group passes 90% of its limit

ORCHID_NO, CROWN_NO = 137, 42
ORCHID_LIMIT = 1_500_000_000
ORCHID_UTILISATION = [0.620, 0.628, 0.636, 0.644, 0.652, 0.660, 0.668, 0.676, 0.684, 0.690, 0.695, 0.700,
                      0.710, 0.740, 0.790, 0.840, 0.890, 0.940]
ORCHID_FX_SHARE = 0.25     # share of the Apr→Sep 2026 rise explained by FX (Sep balances × rate change)
CROWN_LIMIT = 2_400_000_000
CROWN_EXPOSURE = 1_300_000_000

TL, RCF, WCL, PF = "Term Loan", "Revolving Credit Facility", "Working Capital Loan", "Project Finance Loan"

ORCHID_SUBS = [
    # (English name, local name, country, sector, growth driver?, [(product, currency, weight)])
    ("Orchid Holdings Pte. Ltd.", "", "SG", "Infrastructure", False, [(TL, "USD", 4.0), (RCF, "SGD", 4.0)]),
    ("Orchid Logistics Pte. Ltd.", "", "SG", "Shipping & Logistics", False, [(TL, "SGD", 5.0), (RCF, "USD", 2.0)]),
    ("Orchid Capital Pte. Ltd.", "", "SG", "Financial Institutions", False, [(RCF, "USD", 4.0)]),
    ("Orchid Logistics (Thailand) Co., Ltd.", "บริษัท ออร์คิด โลจิสติกส์ (ประเทศไทย) จำกัด", "TH",
     "Shipping & Logistics", False, [(TL, "THB", 6.0), (WCL, "THB", 2.0)]),
    ("Orchid Industries (Thailand) Co., Ltd.", "บริษัท ออร์คิด อินดัสทรีส์ (ประเทศไทย) จำกัด", "TH",
     "Manufacturing", False, [(TL, "THB", 6.0), (RCF, "USD", 1.5)]),
    ("Orchid Consumer (Thailand) Co., Ltd.", "บริษัท ออร์คิด คอนซูมเมอร์ (ประเทศไทย) จำกัด", "TH",
     "Consumer Goods", False, [(WCL, "THB", 4.0)]),
    ("Orchid Land (Thailand) Co., Ltd.", "บริษัท ออร์คิด พร็อพเพอร์ตี้ (ประเทศไทย) จำกัด", "TH",
     "Real Estate", True, [(TL, "THB", 7.0), (TL, "THB", 4.0)]),
    ("Orchid Land Vietnam JSC", "Công ty Cổ phần Bất động sản Hoa Lan Việt Nam", "VN",
     "Real Estate", True, [(TL, "VND", 8.0), (TL, "USD", 2.0)]),
    ("Orchid Energy Vietnam JSC", "Công ty Cổ phần Năng lượng Hoa Lan Việt Nam", "VN",
     "Energy", False, [(TL, "VND", 5.0)]),
    ("PT Orchid Energi Indonesia", "", "ID", "Energy", True, [(TL, "IDR", 9.0), (TL, "USD", 2.0)]),
    ("PT Orchid Properti Indonesia", "", "ID", "Real Estate", True, [(TL, "IDR", 8.0)]),
    ("PT Orchid Logistik Indonesia", "", "ID", "Shipping & Logistics", False, [(RCF, "IDR", 4.0)]),
    ("Orchid Industries India Private Limited", "", "IN", "Manufacturing", False,
     [(TL, "INR", 5.0), (RCF, "USD", 1.0)]),
    ("Orchid Infrastructure India Private Limited", "", "IN", "Infrastructure", False, [(TL, "INR", 5.0)]),
]

CROWN_SUBS = [
    ("Pacific Crown Holdings Limited", "太平洋皇冠控股有限公司", "HK", "Financial Institutions",
     [(RCF, "USD", 10)]),
    ("Pacific Crown Land Limited", "太平洋皇冠置業有限公司", "HK", "Real Estate", [(TL, "HKD", 25)]),
    ("Pacific Crown Telecom Limited", "太平洋皇冠電訊有限公司", "HK", "Telecommunications", [(TL, "HKD", 12)]),
    ("Pacific Crown Energy (China) Co., Ltd.", "太平洋皇冠能源有限公司", "CN", "Energy", [(PF, "CNY", 15)]),
    ("Pacific Crown Logistics (China) Co., Ltd.", "太平洋皇冠物流有限公司", "CN", "Shipping & Logistics",
     [(TL, "CNY", 10)]),
    ("Pacific Crown Capital Pte. Ltd.", "", "SG", "Financial Institutions", [(RCF, "USD", 8)]),
    ("Pacific Crown Resources (Australia) Pty Ltd", "", "AU", "Metals & Mining", [(TL, "AUD", 12)]),
    ("Pacific Crown Technologies (Taiwan) Co., Ltd.", "太平洋皇冠科技有限公司", "TW", "Technology",
     [(RCF, "TWD", 8)]),
]

# Planted data-quality problems.
N_NEGATIVE_COMMITMENT = 12
N_MATURITY_BEFORE_START = 9
N_UNKNOWN_COUNTERPARTY = 15
CURRENCY_VARIANT_RATE = 0.04
COUNTERPARTY_DUPLICATE_RATE = 0.02
RERUN_BRANCH, RERUN_SYSTEM = "BKK", "TRADE_FIN"   # this September 2026 extract was loaded twice
CURRENCY_VARIANTS = {"USD": ["usd", "USD ", "US$", "Usd"], "SGD": ["S$", "sgd"], "CNY": ["RMB"],
                     "JPY": ["jpy", "JPY "]}

# Planted: trade-finance arrears at the Ho Chi Minh City branch creep up from April 2026.
# Target share of the branch's trade-finance exposure that is more than 30 days past due.
DPD_CREEP_BRANCH = "SGN"
DPD_CREEP_TARGET_PCT = [1.0] * APR26 + [2.0, 3.5, 5.0, 6.5, 8.0, 9.5]


def is_rerun(fac, i):
    return i == SEP26 and fac["booking_branch"] == RERUN_BRANCH and fac["source_system"] == RERUN_SYSTEM

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def num_str(x):
    """Plain decimal with 8 significant figures (no scientific notation)."""
    if x == 0:
        return "0"
    decimals = max(0, 8 - 1 - int(math.floor(math.log10(abs(x)))))
    return f"{x:.{decimals}f}".rstrip("0").rstrip(".")


def ceil_sig(x, sig=3):
    """Round a positive number up to `sig` significant figures."""
    step = 10 ** (int(math.floor(math.log10(x))) - sig + 1)
    return int(math.ceil(x / step - 1e-9) * step) if step >= 1 else math.ceil(x)


def weighted(rng, weights):
    keys = list(weights)
    return rng.choices(keys, weights=[weights[k] for k in keys])[0]


def poisson(rng, lam):
    limit, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= limit:
            return k
        k += 1


def is_katakana(text):
    return bool(text) and "゠" <= text[0] <= "ヿ"


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def build_fx(rng):
    usd = {}
    for ccy, (base, vol) in CURRENCIES.items():
        path = [base]
        for _ in range(1, len(MONTHS)):
            path.append(path[-1] * math.exp(rng.gauss(0, vol)))
        if ccy in PLANTED_FX_MOVE:
            start, steps = path[APR26], SEP26 - APR26
            for k in range(1, steps + 1):
                wobble = 1.0 if k == steps else 1 + rng.gauss(0, 0.003)
                path[APR26 + k] = start * (1 + PLANTED_FX_MOVE[ccy] * k / steps) * wobble
        # Store exactly what the CSV will contain.
        usd[ccy] = [float(num_str(x)) for x in path]
    jpy = {c: [float(num_str(usd[c][i] / usd["JPY"][i])) for i in range(len(MONTHS))] for c in usd}
    return usd, jpy


class Builder:
    def __init__(self, rng, usd):
        self.rng, self.usd = rng, usd
        self.groups, self.counterparties, self.facilities = [], [], []
        self.cp_names = set()
        self.next_cp, self.next_fac = 1, 1
        local_combos = [(s, x) for s in LOCAL_STEMS for x in GROUP_SUFFIXES]
        other_combos = [(s, x) for s in OTHER_STEMS for x in GROUP_SUFFIXES]
        rng.shuffle(local_combos)
        rng.shuffle(other_combos)
        self.local_combos, self.other_combos = local_combos, other_combos

    # -- naming -------------------------------------------------------------

    def group_name(self, home):
        if COUNTRIES[home][4] and self.local_combos and self.rng.random() < 0.85:
            return self.local_combos.pop()
        return self.other_combos.pop()

    def counterparty_names(self, stem, sector, country, foreign):
        qualifiers = [""] + NAME_QUALIFIERS + [f"No. {n}" for n in range(2, 100)]
        for qualifier in qualifiers:
            english = english_name(stem, sector, country, foreign, qualifier)
            if english not in self.cp_names:
                self.cp_names.add(english)
                return english, (local_name(stem, sector, country, foreign) if not qualifier else "")
        raise RuntimeError(f"Ran out of names for {stem}/{sector}/{country}")

    # -- entities -----------------------------------------------------------

    def add_counterparty(self, group, name, local, country, sector, rating, rm):
        rng = self.rng
        cp = {
            "counterparty_id": f"CP{self.next_cp:05d}",
            "counterparty_name": name,
            "counterparty_name_local": local,
            "parent_group_id": group["group_id"],
            "country_code": country,
            "sector": sector,
            "internal_rating": max(1, min(10, rating)),
            "relationship_manager": rm,
            "onboarded_on": dt.date(rng.randint(2008, 2024), rng.randint(1, 12), rng.randint(1, 28)),
            "last_updated_at": dt.datetime(2026, rng.randint(1, 9), rng.randint(1, 28),
                                           rng.randint(0, 23), rng.randint(0, 59), rng.randint(0, 59)),
            "group": group,
        }
        self.next_cp += 1
        self.counterparties.append(cp)
        return cp

    def add_facility(self, cp, product, currency, branch, start, maturity):
        rng = self.rng
        sector = cp["sector"]
        secured = rng.random() < (0.85 if sector == "Real Estate" else 0.40)
        fac = {
            "facility_id": f"FAC{self.next_fac:06d}",
            "cp": cp,
            "product_type": product,
            "source_system": PRODUCTS[product][0],
            "currency": currency,
            "start_date": start,
            "maturity_date": maturity,
            "booking_branch": branch,
            "secured": secured,
            "collateral_type": COLLATERAL.get(sector, rng.choice(["Receivables", "Inventory", "Shares"]))
            if secured else None,
            "drawn": {},       # month index -> drawn amount in facility currency
            "dpd": {},         # month index -> days past due
            "issue": None,     # quarantine reason, if planted
        }
        self.next_fac += 1
        self.facilities.append(fac)
        return fac

    def booking_branch(self, country):
        if country == "SG":
            return "SIN"
        if country == "JP":
            return "SIN" if self.rng.random() < 0.6 else "HKG"
        r = self.rng.random()
        return COUNTRIES[country][2] if r < 0.70 else ("SIN" if r < 0.95 else "HKG")

    def facility_currency(self, country):
        local = COUNTRIES[country][1]
        usd_share = {"JP": 0.45, "HK": 0.60}.get(country, 0.50)
        return "USD" if self.rng.random() < usd_share else local

    # -- ordinary groups ----------------------------------------------------

    def standard_group(self, gid, big_ticket):
        rng = self.rng
        home = weighted(rng, COUNTRY_WEIGHTS)
        sector = weighted(rng, {s: v[2] for s, v in SECTORS.items()})
        stem, suffix = self.group_name(home)
        if big_ticket:
            size, target = 1, rng.uniform(230e6, 420e6)
        else:
            r = rng.random()
            size = 1 if r < 0.35 else rng.randint(2, 5) if r < 0.70 else rng.randint(6, 15) if r < 0.92 \
                else rng.randint(16, 30)
            target = min(1.0e9, max(2e6, math.exp(rng.gauss(math.log(35e6), 1.05))))
            if size == 1:
                target *= 0.5
        group = {"group_id": gid, "group_name": f"{stem} {suffix}", "home_country": home,
                 "primary_sector": sector, "kind": "big_ticket" if big_ticket else "standard",
                 "target_usd": target, "u_max": rng.uniform(0.25, 0.85)}
        self.groups.append(group)

        base_rating = rng.choices(range(2, 9), weights=[4, 10, 18, 22, 20, 14, 8])[0]
        rm = f"RM{rng.randint(1, 60):03d}"
        shares = [rng.random() ** 2 + 0.05 for _ in range(size)]
        for j in range(size):
            country = home if (j == 0 or rng.random() < 0.6) else weighted(rng, COUNTRY_WEIGHTS)
            cp_sector = sector if (j == 0 or rng.random() < 0.6) else \
                weighted(rng, {s: v[2] for s, v in SECTORS.items()})
            name, local = self.counterparty_names(stem, cp_sector, country, country != home)
            cp = self.add_counterparty(group, name, local, country, cp_sector,
                                       base_rating + rng.choice([-1, 0, 0, 1]), rm)
            n_fac = 1 if big_ticket else min(7, 1 + poisson(rng, 1.6))
            weights = [rng.random() + 0.2 for _ in range(n_fac)]
            for k in range(n_fac):
                product = weighted(rng, PRODUCT_MIX[SECTORS[cp_sector][3]])
                if big_ticket:
                    product = rng.choice([TL, PF if cp_sector in ("Energy", "Utilities", "Infrastructure") else RCF])
                _, tmin, tmax, profile, expected = PRODUCTS[product]
                tenor = rng.randint(tmin, tmax)
                earliest = add_months(MONTHS[0], -tenor + 1)
                latest = dt.date(2026, 8, 31)
                start = earliest + dt.timedelta(days=rng.randint(0, (latest - earliest).days))
                maturity = add_months(start, tenor)
                currency = self.facility_currency(country)
                fac = self.add_facility(cp, product, currency, self.booking_branch(country), start, maturity)
                target_usd = target * shares[j] / math.fsum(shares) * weights[k] / math.fsum(weights)
                fac["committed_raw"] = target_usd / expected / self.usd[currency][0]
                fac["util"] = utilisation_path(rng, profile, start, maturity)

    def scale_standard_groups(self):
        by_group = defaultdict(list)
        for fac in self.facilities:
            if "util" in fac:
                by_group[fac["cp"]["group"]["group_id"]].append(fac)
        for group in self.groups:
            facs = by_group.get(group["group_id"], [])
            exposure = [0.0] * len(MONTHS)
            for fac in facs:
                for i, u in fac["util"].items():
                    exposure[i] += fac["committed_raw"] * u * self.usd[fac["currency"]][i]
            live = [e for e in exposure if e > 0]
            if not live:
                continue
            scale = group["target_usd"] / (math.fsum(live) / len(live))
            scale = min(scale, GROUP_CAP_USD / max(exposure))
            for fac in facs:
                committed = fac["committed_raw"] * scale
                fac["drawn"] = {i: round(committed * u, 2) for i, u in fac["util"].items()}
                fac["committed_amount"] = ceil_sig(committed)
                del fac["util"], fac["committed_raw"]

    # -- planted groups -----------------------------------------------------

    def orchid_group(self, gid):
        rng, usd = self.rng, self.usd
        group = {"group_id": gid, "group_name": "Orchid Group", "home_country": "SG",
                 "primary_sector": "Infrastructure", "kind": "orchid", "limit": ORCHID_LIMIT}
        self.groups.append(group)
        lines = []
        for name, local, country, sector, growth, facs in ORCHID_SUBS:
            self.cp_names.add(name)
            cp = self.add_counterparty(group, name, local, country, sector, rng.choice([4, 4, 5, 5, 6]), "RM014")
            cp["growth_driver"] = growth
            for product, currency, weight in facs:
                start = dt.date(rng.randint(2023, 2024), rng.randint(1, 12), rng.randint(1, 28))
                maturity = add_months(start, rng.randint(54, 84))
                branch = "BOM" if country == "IN" else "SIN"
                lines.append((self.add_facility(cp, product, currency, branch, start, maturity), weight, growth))

        target = [ORCHID_LIMIT * u for u in ORCHID_UTILISATION]
        total_w = math.fsum(w for _, w, _ in lines)
        fx_gain = {c: 1 - usd[c][APR26] / usd[c][SEP26] for c in usd}
        non_usd = [(f, w) for f, w, _ in lines if f["currency"] != "USD"]
        n0 = math.fsum(w for _, w in non_usd) / total_w
        fx_ratio = ORCHID_FX_SHARE * (target[SEP26] - target[APR26]) / target[SEP26]
        alpha = fx_ratio / math.fsum(w / total_w * fx_gain[f["currency"]] for f, w in non_usd)
        beta = (1 - alpha * n0) / (1 - n0)
        assert beta > 0, "Orchid currency mix cannot deliver the planted FX share"

        sep, apr = {}, {}
        for fac, w, growth in lines:
            share = w / total_w * (alpha if fac["currency"] != "USD" else beta)
            sep[fac["facility_id"]] = target[SEP26] * share / usd[fac["currency"]][SEP26]
            apr[fac["facility_id"]] = sep[fac["facility_id"]] * (rng.uniform(0.50, 0.60) if growth
                                                                  else rng.uniform(0.93, 1.01))
        self._rescale(lines, apr, APR26, target[APR26])

        drawn = {APR26: apr, SEP26: sep}
        for i in range(APR26 + 1, SEP26):
            t = ((i - APR26) / (SEP26 - APR26)) ** 1.15
            raw = {fid: apr[fid] + (sep[fid] - apr[fid]) * t for fid in sep}
            drawn[i] = self._rescale(lines, raw, i, target[i])
        for i in range(APR26):
            raw = {fid: apr[fid] * (0.86 + 0.14 * i / APR26) * (1 + rng.gauss(0, 0.01)) for fid in sep}
            drawn[i] = self._rescale(lines, raw, i, target[i])

        for fac, _, _ in lines:
            fac["drawn"] = {i: round(drawn[i][fac["facility_id"]], 2) for i in range(len(MONTHS))}
            fac["committed_amount"] = ceil_sig(max(fac["drawn"].values()) * rng.uniform(1.08, 1.30))
            fac["dpd"] = {i: 0 for i in range(len(MONTHS))}

    def _rescale(self, lines, amounts, i, target_usd):
        total = math.fsum(amounts[f["facility_id"]] * self.usd[f["currency"]][i] for f, _, _ in lines)
        factor = target_usd / total
        for fid in amounts:
            amounts[fid] *= factor
        return amounts

    def crown_group(self, gid):
        rng, usd = self.rng, self.usd
        group = {"group_id": gid, "group_name": "Pacific Crown Holdings", "home_country": "HK",
                 "primary_sector": "Financial Institutions", "kind": "crown", "limit": CROWN_LIMIT}
        self.groups.append(group)
        branch_of = {"HK": "HKG", "CN": "SHA", "SG": "SIN", "AU": "SYD", "TW": "TPE"}
        lines = []
        for name, local, country, sector, facs in CROWN_SUBS:
            self.cp_names.add(name)
            cp = self.add_counterparty(group, name, local, country, sector, rng.choice([3, 3, 4]), "RM007")
            for product, currency, weight in facs:
                start = dt.date(rng.randint(2022, 2024), rng.randint(1, 12), rng.randint(1, 28))
                maturity = add_months(start, rng.randint(60, 120))
                lines.append((self.add_facility(cp, product, currency, branch_of[country], start, maturity),
                              weight))
        total_w = math.fsum(w for _, w in lines)
        for i in range(len(MONTHS)):
            month_total = CROWN_EXPOSURE * (1 + rng.gauss(0, 0.008))
            for fac, w in lines:
                fac["drawn"][i] = round(month_total * w / total_w / usd[fac["currency"]][i], 2)
        for fac, _ in lines:
            fac["committed_amount"] = ceil_sig(max(fac["drawn"].values()) * rng.uniform(1.15, 1.40))
            fac["dpd"] = {i: 0 for i in range(len(MONTHS))}

    # -- days past due ------------------------------------------------------

    def assign_dpd(self):
        """Days past due: background noise everywhere, plus the planted creep at the SGN trade-finance book."""
        rng, usd = self.rng, self.usd
        creeping = [f for f in self.facilities if not f["dpd"] and f["issue"] is None
                    and f["booking_branch"] == DPD_CREEP_BRANCH and f["source_system"] == "TRADE_FIN"]
        creeping_ids = {f["facility_id"] for f in creeping}
        for fac in self.facilities:
            if fac["dpd"] or fac["facility_id"] in creeping_ids:
                continue
            for i in fac["drawn"]:
                r = rng.random()
                fac["dpd"][i] = (rng.randint(31, 89) if r < 0.008 else rng.randint(91, 150) if r < 0.010
                                 else rng.randint(1, 30) if r < 0.045 else 0)

        # Greedy fill: each month, mark facilities as 30+ days late until the target share is reached.
        # Before April 2026 arrears cure within two months; from April 2026 they stay late.
        order = sorted(creeping, key=lambda f: f["facility_id"])
        if not order:
            return
        rng.shuffle(order)
        late_since = {}
        for i in range(len(MONTHS)):
            live = [f for f in creeping if i in f["drawn"]]
            value = {f["facility_id"]: f["drawn"][i] * usd[f["currency"]][i] for f in live}
            total = math.fsum(value.values())
            if i < APR26:
                late_since = {fid: m for fid, m in late_since.items() if i - m < 2}
            late = {fid for fid in value if fid in late_since}
            share = 100 * math.fsum(value[fid] for fid in late) / total if total else 0.0
            if i < APR26:
                # Background arrears: cure the largest ones if balance swings push the share over target.
                for fid in sorted(late, key=lambda f: -value[f]):
                    if share <= DPD_CREEP_TARGET_PCT[i] + 0.75:
                        break
                    late.discard(fid)
                    del late_since[fid]
                    share -= 100 * value[fid] / total
            # Prefer typical-sized facilities so the share by count stays close to the share by exposure.
            offset = (7 * i) % len(order)
            mean = total / len(value) if value else 0.0
            rotated = order[offset:] + order[:offset]
            for fac in sorted(rotated, key=lambda f: abs(value.get(f["facility_id"], 0.0) - mean)):
                fid = fac["facility_id"]
                if share >= DPD_CREEP_TARGET_PCT[i]:
                    break
                if fid not in value or fid in late:
                    continue
                add = 100 * value[fid] / total
                if share + add > DPD_CREEP_TARGET_PCT[i] + 0.75:
                    continue
                late_since[fid], share = i, share + add
                late.add(fid)
            for fac in live:
                fid = fac["facility_id"]
                if fid in late:
                    fac["dpd"][i] = min(180, 31 + 30 * (i - late_since[fid]) + rng.randint(0, 9))
                else:
                    fac["dpd"][i] = rng.randint(1, 30) if rng.random() < 0.04 else 0

    def build(self):
        rng = self.rng
        candidates = [n for n in range(1, N_GROUPS + 1) if n not in (ORCHID_NO, CROWN_NO)]
        big_ticket = set(rng.sample(candidates, N_BIG_TICKET))
        for n in range(1, N_GROUPS + 1):
            gid = f"GRP{n:04d}"
            if n == ORCHID_NO:
                self.orchid_group(gid)
            elif n == CROWN_NO:
                self.crown_group(gid)
            else:
                self.standard_group(gid, n in big_ticket)
        self.scale_standard_groups()


def english_name(stem, sector, country, foreign, qualifier=""):
    word = SECTORS[sector][0] + (f" {qualifier}" if qualifier else "")
    country_name, _, _, suffix, _ = COUNTRIES[country]
    if country == "ID":
        word = SECTORS[sector][1] + (f" {qualifier}" if qualifier else "")
        return f"PT {stem} {word}" + (" Indonesia" if foreign else "")
    if country == "VN":
        return f"{stem} {word} {'Vietnam ' if foreign else ''}JSC"
    if country == "IN":
        return f"{stem} {word} {'India ' if foreign else ''}{suffix}"
    return f"{stem} {word}{f' ({country_name})' if foreign else ''} {suffix}"


def local_name(stem, sector, country, foreign):
    script = COUNTRIES[country][4]
    if script is None or stem not in LOCAL_STEMS:
        return ""
    idx = SCRIPT_INDEX[script]
    stem_local, sector_local = LOCAL_STEMS[stem][idx], SECTOR_LOCAL[sector][idx]
    if script == "ja":
        return f"{stem_local}{'・' if is_katakana(sector_local) else ''}{sector_local}株式会社"
    if script in ("zh", "zht"):
        return f"{stem_local}{sector_local}有限公司"
    if script == "th":
        return f"บริษัท {stem_local} {sector_local}{' (ประเทศไทย)' if foreign else ''} จำกัด"
    if script == "ko":
        return f"{stem_local} {sector_local} 주식회사"
    return f"Công ty Cổ phần {sector_local} {stem_local}{' Việt Nam' if foreign else ''}"


def utilisation_path(rng, profile, start, maturity):
    path, mean = {}, rng.uniform(0.3, 0.8)
    level, tenor = mean, max(1, months_between(start, maturity))
    for i, d in enumerate(MONTHS):
        if d < start or d > maturity:
            continue
        age = months_between(start, d)
        if profile == "amortising":
            path[i] = max(0.25, 0.98 - 0.7 * age / tenor) * rng.uniform(0.97, 1.0)
        elif profile == "construction":
            path[i] = min(0.97, 0.15 + age / 30.0)
        elif profile == "revolving":
            level = min(0.97, max(0.05, mean + 0.7 * (level - mean) + rng.gauss(0, 0.08)))
            path[i] = level
        else:
            path[i] = min(1.0, max(0.1, rng.gauss(0.6, 0.2)))
    return path


# ---------------------------------------------------------------------------
# Planted data-quality problems
# ---------------------------------------------------------------------------


def plant_issues(rng, builder):
    eligible = [f for f in builder.facilities
                if f["cp"]["group"]["kind"] == "standard" and len(f["drawn"]) >= 3]
    picks = rng.sample(eligible, N_NEGATIVE_COMMITMENT + N_MATURITY_BEFORE_START + N_UNKNOWN_COUNTERPARTY)
    for n, fac in enumerate(picks):
        if n < N_NEGATIVE_COMMITMENT:
            fac["issue"] = "NON_POSITIVE_COMMITMENT"
        elif n < N_NEGATIVE_COMMITMENT + N_MATURITY_BEFORE_START:
            fac["issue"] = "MATURITY_BEFORE_START"
        else:
            fac["issue"] = "UNKNOWN_COUNTERPARTY"
            fac["orphan_id"] = f"CP{90000 + n:05d}"

    variants = {}
    for fac in builder.facilities:
        if fac["cp"]["group"]["kind"] == "standard" and fac["issue"] is None \
                and rng.random() < CURRENCY_VARIANT_RATE:
            options = CURRENCY_VARIANTS.get(fac["currency"], [fac["currency"].lower(), fac["currency"] + " "])
            fac["currency_raw"] = rng.choice(options)
            variants[fac["currency_raw"]] = variants.get(fac["currency_raw"], 0) + 1

    standard_cps = [c for c in builder.counterparties if c["group"]["kind"] == "standard"]
    duplicated = rng.sample(standard_cps, int(len(standard_cps) * COUNTERPARTY_DUPLICATE_RATE))
    for cp in duplicated:
        cp["duplicated"] = True
    return variants


def set_limits(builder, usd):
    """Choose group limits so that only Orchid ever passes 90% utilisation, even on dirty data."""
    clean = defaultdict(lambda: [0.0] * len(MONTHS))
    dirty = defaultdict(lambda: [0.0] * len(MONTHS))
    for fac in builder.facilities:
        gid = fac["cp"]["group"]["group_id"]
        copies = 2 if fac["cp"].get("duplicated") else 1
        for i, amount in fac["drawn"].items():
            value = amount * usd[fac["currency"]][i]
            if fac["issue"] is None:
                clean[gid][i] += value
            if fac["issue"] != "UNKNOWN_COUNTERPARTY":
                dirty[gid][i] += value * copies * (2 if is_rerun(fac, i) else 1)
    for group in builder.groups:
        if group["kind"] in ("orchid", "crown"):
            continue
        gid = group["group_id"]
        need = max(max(c / group["u_max"], d / DIRTY_UTILISATION_CAP) for c, d in zip(clean[gid], dirty[gid]))
        group["limit"] = ceil_sig(max(need, 5e6), 2)


# ---------------------------------------------------------------------------
# Answer key (computed from the same values that are written to the files)
# ---------------------------------------------------------------------------


def answer_key(builder, usd, jpy, variants, seed):
    groups = {g["group_id"]: g for g in builder.groups}
    valid = [f for f in builder.facilities if f["issue"] is None]

    def exposure(fac, i):
        return fac["drawn"][i] * usd[fac["currency"]][i]

    group_exp = defaultdict(lambda: [0.0] * len(MONTHS))
    for fac in valid:
        for i in fac["drawn"]:
            group_exp[fac["cp"]["group"]["group_id"]][i] += exposure(fac, i)

    def ranked(i, n=10):
        rows = sorted(((gid, e[i]) for gid, e in group_exp.items()), key=lambda r: -r[1])[:n]
        return [{"rank": r + 1, "group_id": gid, "group_name": groups[gid]["group_name"],
                 "exposure_usd": round(e, 2), "utilisation_pct": round(100 * e / groups[gid]["limit"], 4)}
                for r, (gid, e) in enumerate(rows)]

    def over_90(i):
        return sorted(groups[gid]["group_name"] for gid, e in group_exp.items()
                      if e[i] > 0.9 * groups[gid]["limit"])

    orchid = next(g for g in builder.groups if g["kind"] == "orchid")
    orchid_facs = [f for f in valid if f["cp"]["group"] is orchid]
    e_apr = math.fsum(exposure(f, APR26) for f in orchid_facs)
    e_sep = math.fsum(exposure(f, SEP26) for f in orchid_facs)
    volume = math.fsum((f["drawn"][SEP26] - f["drawn"][APR26]) * usd[f["currency"]][APR26] for f in orchid_facs)
    fx = math.fsum(f["drawn"][SEP26] * (usd[f["currency"]][SEP26] - usd[f["currency"]][APR26]) for f in orchid_facs)
    fx_alt = math.fsum(f["drawn"][APR26] * (usd[f["currency"]][SEP26] - usd[f["currency"]][APR26])
                 for f in orchid_facs)

    cp_sep = defaultdict(float)
    for fac in valid:
        if SEP26 in fac["drawn"]:
            cp_sep[fac["cp"]["counterparty_id"]] += exposure(fac, SEP26)
    cp_rank = {cid: r + 1 for r, (cid, _) in enumerate(sorted(cp_sep.items(), key=lambda kv: -kv[1]))}
    subs = []
    for cp in [c for c in builder.counterparties if c["group"] is orchid]:
        facs = [f for f in orchid_facs if f["cp"] is cp]
        subs.append({
            "counterparty_id": cp["counterparty_id"], "counterparty_name": cp["counterparty_name"],
            "country_code": cp["country_code"], "growth_driver": cp["growth_driver"],
            "exposure_usd_2026_09_30": round(math.fsum(exposure(f, SEP26) for f in facs), 2),
            "volume_change_usd_apr_to_sep_2026": round(math.fsum(
                (f["drawn"][SEP26] - f["drawn"][APR26]) * usd[f["currency"]][APR26] for f in facs), 2),
            "rank_among_all_counterparties_2026_09_30": cp_rank[cp["counterparty_id"]],
        })
    subs.sort(key=lambda s: -s["exposure_usd_2026_09_30"])
    drivers_volume = math.fsum(s["volume_change_usd_apr_to_sep_2026"] for s in subs if s["growth_driver"])

    rerun_usd = math.fsum(exposure(f, SEP26) for f in valid if SEP26 in f["drawn"] and is_rerun(f, SEP26))
    fanout_usd = math.fsum(exposure(f, SEP26) for f in valid if SEP26 in f["drawn"] and f["cp"].get("duplicated"))
    quarantined_usd = math.fsum(exposure(f, SEP26) for f in builder.facilities
                          if f["issue"] and SEP26 in f["drawn"])
    tradefin_usd = math.fsum(exposure(f, SEP26) for f in valid
                       if SEP26 in f["drawn"] and f["source_system"] == "TRADE_FIN")

    creep = []
    for i, d in enumerate(MONTHS):
        facs = [f for f in valid if i in f["drawn"] and f["booking_branch"] == DPD_CREEP_BRANCH
                and f["source_system"] == "TRADE_FIN"]
        total = math.fsum(exposure(f, i) for f in facs)
        late = math.fsum(exposure(f, i) for f in facs if f["dpd"][i] > 30)
        creep.append({"as_of_date": d.isoformat(), "share_of_exposure_over_30_dpd_pct":
                      round(100 * late / total, 4) if total else None})
    port_total = math.fsum(exposure(f, SEP26) for f in valid if SEP26 in f["drawn"])
    port_late = math.fsum(exposure(f, SEP26) for f in valid if SEP26 in f["drawn"] and f["dpd"][SEP26] > 30)

    def max_util(exposures, kinds_excluded=("orchid",)):
        return round(100 * max(e / groups[gid]["limit"] for gid, es in exposures.items()
                               if groups[gid]["kind"] not in kinds_excluded for e in es), 4)

    balance_rows = sum(len(f["drawn"]) for f in builder.facilities)
    rerun_rows = sum(1 for f in builder.facilities if SEP26 in f["drawn"] and is_rerun(f, SEP26))
    issues = defaultdict(int)
    for f in builder.facilities:
        if f["issue"]:
            issues[f["issue"]] += 1

    return {
        "generated_with": {"script": "generate_data.py", "seed": seed},
        "reporting_dates": {"calendar_h1_2026_end": MONTHS[JUN26].isoformat(),
                            "fy2026_h1_end": MONTHS[SEP26].isoformat(),
                            "note": "Fiscal year runs April-March; FY2026 = 1 Apr 2026 to 31 Mar 2027."},
        "row_counts": {
            "counterparties_unique": len(builder.counterparties),
            "counterparty_rows_in_file": len(builder.counterparties)
            + sum(1 for c in builder.counterparties if c.get("duplicated")),
            "facilities": len(builder.facilities),
            "balance_rows_unique": balance_rows,
            "balance_rows_in_file": balance_rows + rerun_rows,
            "groups": len(builder.groups),
        },
        "planted_data_quality": {
            "counterparty_duplicate_rows": sum(1 for c in builder.counterparties if c.get("duplicated")),
            "currency_codes_needing_cleanup": dict(sorted(variants.items())),
            "quarantine_counts": dict(sorted(issues.items())),
            "balance_rerun_duplicate_rows": rerun_rows,
            "balance_rerun_extract": f"{RERUN_BRANCH} {RERUN_SYSTEM} extract for {MONTHS[SEP26].isoformat()}",
            "balance_rows_with_dd_mm_yyyy_dates": sum(len(f["drawn"]) for f in builder.facilities
                                                      if f["source_system"] == "TRADE_FIN") + rerun_rows,
            "facility_rows_with_dd_mm_yyyy_dates": sum(1 for f in builder.facilities
                                                       if f["source_system"] == "TRADE_FIN"),
            "fx_months_where_rate_date_is_not_month_end": [
                d.isoformat() for d in MONTHS if last_business_day(d) != d],
        },
        "totals_2026_09_30": {
            "exposure_usd": round(port_total, 2),
            "exposure_jpy": round(math.fsum(f["drawn"][SEP26] * jpy[f["currency"]][SEP26] for f in valid
                                      if SEP26 in f["drawn"]), 2),
            "overstatement_if_rerun_not_deduplicated_usd": round(rerun_usd, 2),
            "overstatement_if_counterparties_not_deduplicated_usd": round(fanout_usd, 2),
            "excluded_by_quarantine_usd": round(quarantined_usd, 2),
            "lost_if_only_iso_dates_parsed_usd": round(tradefin_usd, 2),
        },
        "groups": {
            "top10_by_exposure_2026_06_30": ranked(JUN26),
            "top10_by_exposure_2026_09_30": ranked(SEP26),
            "over_90pct_of_limit_2026_06_30": over_90(JUN26),
            "over_90pct_of_limit_2026_09_30": over_90(SEP26),
            "max_utilisation_pct_excluding_orchid_any_month": max_util(group_exp),
        },
        "orchid_group": {
            "group_id": orchid["group_id"],
            "group_limit_usd": orchid["limit"],
            "monthly": [{"as_of_date": d.isoformat(), "exposure_usd": round(group_exp[orchid["group_id"]][i], 2),
                         "utilisation_pct": round(100 * group_exp[orchid["group_id"]][i] / orchid["limit"], 4)}
                        for i, d in enumerate(MONTHS)],
            "first_month_over_90pct": next(d.isoformat() for i, d in enumerate(MONTHS)
                                           if group_exp[orchid["group_id"]][i] > 0.9 * orchid["limit"]),
            "change_apr_to_sep_2026": {
                "exposure_2026_04_30_usd": round(e_apr, 2),
                "exposure_2026_09_30_usd": round(e_sep, 2),
                "total_change_usd": round(e_sep - e_apr, 2),
                "volume_effect_usd": round(volume, 2),
                "fx_effect_usd": round(fx, 2),
                "fx_share_pct": round(100 * fx / (e_sep - e_apr), 4),
                "method": "volume = (Sep balance - Apr balance) x Apr rate; FX = Sep balance x (Sep rate - Apr rate)",
                "alt_fx_effect_on_april_balances_usd": round(fx_alt, 2),
                "alt_fx_share_pct": round(100 * fx_alt / (e_sep - e_apr), 4),
            },
            "growth_drivers_share_of_volume_change_pct": round(100 * drivers_volume / volume, 4),
            "subsidiaries": subs,
        },
        "dpd_creep": {
            "booking_branch": DPD_CREEP_BRANCH,
            "source_system": "TRADE_FIN",
            "by_month": creep,
            "portfolio_share_of_exposure_over_30_dpd_pct_2026_09_30": round(100 * port_late / port_total, 4),
        },
    }


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(path, sheet, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return write_xlsx_minimal(path, sheet, header, rows)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFG", [11, 34, 13, 24, 18, 18, 15]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = "#,##0"
        row[5].number_format = "yyyy-mm-dd"
    wb.save(path)


def write_xlsx_minimal(path, sheet, header, rows):
    """Write a single-sheet .xlsx without third-party libraries (dates are written as ISO text)."""
    def ref(c, r):
        letters = ""
        c += 1
        while c:
            c, rem = divmod(c - 1, 26)
            letters = chr(65 + rem) + letters
        return f"{letters}{r}"

    xml_rows = []
    for r, row in enumerate([header] + rows, start=1):
        cells = []
        for c, value in enumerate(row):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cells.append(f'<c r="{ref(c, r)}"><v>{value}</v></c>')
            else:
                text = value.isoformat() if isinstance(value, dt.date) else str(value)
                cells.append(f'<c r="{ref(c, r)}" t="inlineStr"><is><t>{escape(text)}</t></is></c>')
        xml_rows.append(f'<row r="{r}">{"".join(cells)}</row>')
    ns = "http://schemas.openxmlformats.org"
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Types xmlns="{ns}/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>'),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{ns}/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{ns}/officeDocument/2006/relationships/officeDocument" '
            'Target="xl/workbook.xml"/></Relationships>'),
        "xl/workbook.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<workbook xmlns="{ns}/spreadsheetml/2006/main" xmlns:r="{ns}/officeDocument/2006/relationships">'
            f'<sheets><sheet name="{escape(sheet)}" sheetId="1" r:id="rId1"/></sheets></workbook>'),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{ns}/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{ns}/officeDocument/2006/relationships/worksheet" '
            'Target="worksheets/sheet1.xml"/></Relationships>'),
        "xl/worksheets/sheet1.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{ns}/spreadsheetml/2006/main"><sheetData>{"".join(xml_rows)}</sheetData>'
            '</worksheet>'),
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in files.items():
            zf.writestr(name, content)


def write_outputs(root, rng, builder, usd, jpy, key):
    land = os.path.join(root, "landing")
    fac_dir = os.path.join(root, "facilitator")

    # counterparties.csv, with the planted duplicate rows mixed in
    cp_rows = []
    for cp in builder.counterparties:
        cp_rows.append([cp["counterparty_id"], cp["counterparty_name"], cp["counterparty_name_local"],
                        cp["parent_group_id"], cp["country_code"], cp["sector"], cp["internal_rating"],
                        cp["relationship_manager"], cp["onboarded_on"].isoformat(),
                        cp["last_updated_at"].strftime("%Y-%m-%d %H:%M:%S")])
    for cp in [c for c in builder.counterparties if c.get("duplicated")]:
        original = next(r for r in cp_rows if r[0] == cp["counterparty_id"])
        cp_rows.insert(rng.randint(0, len(cp_rows)), list(original))
    write_csv(os.path.join(land, "counterparties", "counterparties.csv"),
              ["counterparty_id", "counterparty_name", "counterparty_name_local", "parent_group_id",
               "country_code", "sector", "internal_rating", "relationship_manager", "onboarded_on",
               "last_updated_at"], cp_rows)

    # facilities.json (JSON Lines)
    path = os.path.join(land, "facilities", "facilities.json")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        for fac in builder.facilities:
            system = fac["source_system"]
            start, maturity = fac["start_date"], fac["maturity_date"]
            if fac["issue"] == "MATURITY_BEFORE_START":
                maturity = start.replace(year=start.year - 1, day=min(start.day, 28))
            committed = fac["committed_amount"]
            if fac["issue"] == "NON_POSITIVE_COMMITMENT":
                committed = -committed
            collateral = None
            if fac["secured"]:
                value_usd = fac["committed_amount"] * usd[fac["currency"]][0] * rng.uniform(0.6, 1.4)
                collateral = {"type": fac["collateral_type"], "value_usd": round(value_usd, -3)}
            record = {
                "facility_id": fac["facility_id"],
                "counterparty_id": fac.get("orphan_id", fac["cp"]["counterparty_id"]),
                "product_type": fac["product_type"],
                "currency": fac.get("currency_raw", fac["currency"]),
                "committed_amount": committed,
                "start_date": fmt_date(start, system),
                "maturity_date": fmt_date(maturity, system),
                "booking_branch": fac["booking_branch"],
                "secured": fac["secured"],
                "collateral": collateral,
                "source_system": system,
            }
            fh.write(json.dumps(record) + "\n")

    # monthly_balances.csv: one extract per month, then the re-run of the BKK September extract
    rows, rerun = [], []
    for i, d in enumerate(MONTHS):
        loaded = dt.datetime.combine(d + dt.timedelta(days=2), dt.time(6, 15))
        for fac in builder.facilities:
            if i not in fac["drawn"]:
                continue
            row = [fmt_date(d, fac["source_system"]), fac["facility_id"], f"{fac['drawn'][i]:.2f}",
                   fac["dpd"][i], fac["source_system"],
                   (loaded + dt.timedelta(minutes=rng.randint(0, 40))).strftime("%Y-%m-%d %H:%M:%S")]
            rows.append(row)
            if is_rerun(fac, i):
                rerun.append(row[:5] + ["2026-10-03 21:40:00"])
    write_csv(os.path.join(land, "monthly_balances", "monthly_balances.csv"),
              ["as_of_date", "facility_id", "drawn_amount", "days_past_due", "source_system", "extract_ts"],
              rows + rerun)

    # fx_rates.csv: one rate per month, dated on the last business day
    fx_rows = [[last_business_day(d).isoformat(), ccy, num_str(usd[ccy][i]), num_str(jpy[ccy][i])]
               for i, d in enumerate(MONTHS) for ccy in CURRENCIES]
    write_csv(os.path.join(land, "fx_rates", "fx_rates.csv"),
              ["rate_date", "currency_code", "usd_rate", "jpy_rate"], fx_rows)

    # group_limits.xlsx (the risk team's spreadsheet) plus a CSV copy for automation
    limit_rows = []
    for g in sorted(builder.groups, key=lambda g: g["group_id"]):
        approved = dt.date(2025, 4, 1) + dt.timedelta(days=rng.randint(0, 420))
        limit_rows.append([g["group_id"], g["group_name"], g["home_country"], g["primary_sector"],
                           int(g["limit"]), approved, rng.choice(CREDIT_OFFICERS)])
    header = ["group_id", "group_name", "home_country", "primary_sector", "group_limit_usd",
              "limit_approved_on", "credit_officer"]
    write_xlsx(os.path.join(land, "group_limits", "group_limits.xlsx"), "group_limits", header, limit_rows)
    write_csv(os.path.join(fac_dir, "group_limits.csv"), header,
              [r[:5] + [r[5].isoformat(), r[6]] for r in limit_rows])

    write_csv(os.path.join(land, "sector_emission_factors", "sector_emission_factors.csv"),
              ["sector", "tco2e_per_usd_million", "note"],
              [[s, f, "Illustrative factor for workshop use only"] for s, f in EMISSION_FACTORS.items()])

    with open(os.path.join(fac_dir, "answer_key.json"), "w", encoding="utf-8") as fh:
        json.dump(key, fh, indent=2, ensure_ascii=False)


def copy_tree(src, dst):
    # File-by-file copy: Unity Catalog volumes don't support the metadata calls copytree makes.
    for dirpath, _, filenames in os.walk(src):
        target = os.path.join(dst, os.path.relpath(dirpath, src))
        os.makedirs(target, exist_ok=True)
        for name in filenames:
            shutil.copyfile(os.path.join(dirpath, name), os.path.join(target, name))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="output", help="output root (local folder or /Volumes/<catalog>/<schema>)")
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args(argv)

    rng = random.Random(args.seed)
    usd, jpy = build_fx(rng)
    builder = Builder(rng, usd)
    builder.build()
    variants = plant_issues(rng, builder)
    builder.assign_dpd()
    set_limits(builder, usd)
    key = answer_key(builder, usd, jpy, variants, args.seed)

    with tempfile.TemporaryDirectory() as tmp:
        write_outputs(tmp, rng, builder, usd, jpy, key)
        copy_tree(tmp, args.out)

    orchid = key["orchid_group"]
    print(f"Wrote {key['row_counts']} to {args.out}")
    print(f"Orchid Group: {orchid['monthly'][APR26]['utilisation_pct']:.1f}% (Apr 2026) -> "
          f"{orchid['monthly'][SEP26]['utilisation_pct']:.1f}% (Sep 2026) of limit; "
          f"FX share of the rise {orchid['change_apr_to_sep_2026']['fx_share_pct']:.1f}%")
    print(f"Groups over 90%: Jun 2026 {key['groups']['over_90pct_of_limit_2026_06_30']}, "
          f"Sep 2026 {key['groups']['over_90pct_of_limit_2026_09_30']}")


if __name__ == "__main__":
    main()
